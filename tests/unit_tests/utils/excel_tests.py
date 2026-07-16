# Licensed to the Apache Software Foundation (ASF) under one
# or more contributor license agreements.  See the NOTICE file
# distributed with this work for additional information
# regarding copyright ownership.  The ASF licenses this file
# to you under the Apache License, Version 2.0 (the
# "License"); you may not use this file except in compliance
# with the License.  You may obtain a copy of the License at
#
#   http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing,
# software distributed under the License is distributed on an
# "AS IS" BASIS, WITHOUT WARRANTIES OR CONDITIONS OF ANY
# KIND, either express or implied.  See the License for the
# specific language governing permissions and limitations
# under the License.

import io
from datetime import datetime, timezone

import pandas as pd
import pytest
from openpyxl import load_workbook
from pandas.api.types import is_numeric_dtype

from superset.utils.core import GenericDataType
from superset.utils.excel import (
    apply_column_types,
    column_number_to_letter,
    df_to_excel,
    get_function_num,
)


def test_timezone_conversion() -> None:
    """
    Test that columns with timezones are converted to a string.
    """
    df = pd.DataFrame({"dt": [datetime(2023, 1, 1, 0, 0, tzinfo=timezone.utc)]})
    apply_column_types(df, [GenericDataType.TEMPORAL])
    contents = df_to_excel(df)
    assert pd.read_excel(contents)["dt"][0] == "2023-01-01 00:00:00+00:00"


def test_quote_formulas() -> None:
    """
    Test that formulas are quoted in Excel.
    """
    df = pd.DataFrame({"formula": ["=SUM(A1:A2)", "normal", "@SUM(A1:A2)"]})
    contents = df_to_excel(df)
    assert pd.read_excel(contents)["formula"].tolist() == [
        "'=SUM(A1:A2)",
        "normal",
        "'@SUM(A1:A2)",
    ]


def test_column_data_types_with_one_numeric_column():
    df = pd.DataFrame(
        {
            "col0": ["123", "1", "2", "3"],
            "col1": ["456", "5.67", "0", ".45"],
            "col2": [
                datetime(2023, 1, 1, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 2, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 3, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 4, 0, 0, tzinfo=timezone.utc),
            ],
            "col3": ["True", "False", "True", "False"],
        }
    )
    coltypes: list[GenericDataType] = [
        GenericDataType.STRING,
        GenericDataType.NUMERIC,
        GenericDataType.TEMPORAL,
        GenericDataType.BOOLEAN,
    ]

    # only col1 should be converted to numeric, according to coltypes definition
    assert not is_numeric_dtype(df["col1"])
    apply_column_types(df, coltypes)
    assert not is_numeric_dtype(df["col0"])
    assert is_numeric_dtype(df["col1"])
    assert not is_numeric_dtype(df["col2"])
    assert not is_numeric_dtype(df["col3"])


def test_column_data_types_with_failing_conversion():
    df = pd.DataFrame(
        {
            "col0": ["123", "1", "2", "3"],
            "col1": ["456", "non_numeric_value", "0", ".45"],
            "col2": [
                datetime(2023, 1, 1, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 2, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 3, 0, 0, tzinfo=timezone.utc),
                datetime(2023, 1, 4, 0, 0, tzinfo=timezone.utc),
            ],
            "col3": ["True", "False", "True", "False"],
        }
    )
    coltypes: list[GenericDataType] = [
        GenericDataType.STRING,
        GenericDataType.NUMERIC,
        GenericDataType.TEMPORAL,
        GenericDataType.BOOLEAN,
    ]

    # should not fail neither convert
    assert not is_numeric_dtype(df["col1"])
    apply_column_types(df, coltypes)
    assert not is_numeric_dtype(df["col0"])
    assert not is_numeric_dtype(df["col1"])
    assert not is_numeric_dtype(df["col2"])
    assert not is_numeric_dtype(df["col3"])


def test_column_data_types_with_large_numeric_values():
    df = pd.DataFrame(
        {
            "big_number": [
                10**14,
                999999999999999,
                10**15 + 1,
                10**16,
                1100108628127863,
                2**54,
            ],
        }
    )
    apply_column_types(df, [GenericDataType.NUMERIC])
    assert df["big_number"].tolist() == [
        100000000000000,
        999999999999999,
        "1000000000000001",
        "10000000000000000",
        "1100108628127863",
        "18014398509481984",
    ]


@pytest.mark.parametrize(
    "column_number,expected",
    [
        (0, "A"),
        (1, "B"),
        (25, "Z"),
        (26, "AA"),
        (27, "AB"),
    ],
)
def test_column_number_to_letter(column_number: int, expected: str) -> None:
    assert column_number_to_letter(column_number) == expected


@pytest.mark.parametrize(
    "aggregate,expected",
    [
        ("SUM", 109),
        ("AVG", 101),
        ("COUNT", 102),
        ("MAX", 104),
        ("MIN", 105),
        ("COUNT_DISTINCT", None),
        (None, None),
    ],
)
def test_get_function_num_ignoring_hidden_rows(aggregate, expected) -> None:
    # default ignore_hidden_rows=True adds 100 to the base SUBTOTAL function code
    assert get_function_num(aggregate) == expected


def test_get_function_num_including_hidden_rows() -> None:
    assert get_function_num("SUM", ignore_hidden_rows=False) == 9


def test_df_to_excel_with_summary_specs() -> None:
    """
    A summary row with a SUBTOTAL formula is appended for each summary spec.
    """
    df = pd.DataFrame({"office": ["a", "b"], "amount": [10, 20]})
    contents = df_to_excel(df, summary_specs=[{"label": "amount", "aggregate": "SUM"}])

    worksheet = load_workbook(io.BytesIO(contents)).active
    cell_values = [cell.value for row in worksheet.iter_rows() for cell in row]

    assert "Summary" in cell_values
    formulas = [
        value
        for value in cell_values
        if isinstance(value, str) and value.startswith("=SUBTOTAL")
    ]
    assert len(formulas) == 1
    # 109 == SUM (9) while ignoring hidden rows (+100)
    assert formulas[0].startswith("=SUBTOTAL(109,")


def test_df_to_excel_without_summary_specs_has_no_formula() -> None:
    df = pd.DataFrame({"office": ["a", "b"], "amount": [10, 20]})
    contents = df_to_excel(df)

    worksheet = load_workbook(io.BytesIO(contents)).active
    cell_values = [cell.value for row in worksheet.iter_rows() for cell in row]

    assert "Summary" not in cell_values
    assert not any(
        isinstance(value, str) and value.startswith("=SUBTOTAL")
        for value in cell_values
    )
